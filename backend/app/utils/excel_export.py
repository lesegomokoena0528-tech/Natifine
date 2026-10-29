from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.report_formatting import display_label, flatten_report_data


class ExcelExporter:

    def generate(

            self,

            filename,

            data

    ):

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "System Report"
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:C1")
        title = sheet["A1"]
        title.value = "NatiFine SA | System Report"
        title.font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
        title.fill = PatternFill("solid", fgColor="12304A")
        title.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34

        sheet.append([])
        sheet.append(["Section", "Metric", "Value"])
        header_fill = PatternFill("solid", fgColor="1F6E8C")
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        section_fills = (
            PatternFill("solid", fgColor="F2F7FA"),
            PatternFill("solid", fgColor="FFFFFF"),
        )
        previous_section = None
        section_index = -1
        for section, metric, value in flatten_report_data(data):
            if section != previous_section:
                section_index += 1
                previous_section = section
            if isinstance(value, (dict, list, tuple)):
                value = str(value)
            elif hasattr(value, "as_tuple"):
                value = float(value)
            sheet.append([
                display_label(section),
                display_label(metric),
                value,
            ])

            row = sheet.max_row
            fill = section_fills[section_index % len(section_fills)]
            for cell in sheet[row]:
                cell.fill = fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, (int, float)) and any(
                word in metric.lower()
                for word in ("revenue", "amount", "total_paid")
            ):
                sheet.cell(row=row, column=3).number_format = '"R" #,##0.00'

        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:C{sheet.max_row}"
        for column, width in {"A": 22, "B": 36, "C": 34}.items():
            sheet.column_dimensions[column].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

        workbook.save(filename)