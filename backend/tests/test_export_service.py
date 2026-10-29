from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.services.export_service import ExportService


def test_csv_export_returns_downloadable_file(app):
    service = ExportService()

    with app.test_request_context():
        response = service.export_csv(
            "system_report",
            {"dashboard": {"users": 3}, "revenue": Decimal("125.50")},
        )
        response.direct_passthrough = False

    assert response.mimetype == "text/csv"
    assert response.headers["Content-Disposition"].startswith(
        "attachment; filename=system_report_"
    )
    assert response.get_data().startswith(b"Section,Metric,Value")
    assert b"Dashboard,Users,3" in response.get_data()


def test_pdf_export_returns_downloadable_file(app):
    service = ExportService()

    with app.test_request_context():
        response = service.export_pdf(
            "system_report",
            {
                "dashboard": {"users": 3},
                "revenue": Decimal("125.50"),
                "generated": "2026-10-29T18:00:00",
            },
        )
        response.direct_passthrough = False

    assert response.mimetype == "application/pdf"
    assert response.headers["Content-Disposition"].startswith(
        "attachment; filename=system_report_"
    )
    pdf = response.get_data()
    assert pdf.startswith(b"%PDF-")
    assert len(pdf) > 1000


def test_excel_export_returns_downloadable_file(app):
    service = ExportService()

    with app.test_request_context():
        response = service.export_excel(
            "system_report",
            {
                "dashboard": {"users": 3, "revenue": Decimal("125.50")},
                "generated": "2026-10-29T18:00:00",
            },
        )
        response.direct_passthrough = False

    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["Content-Disposition"].startswith(
        "attachment; filename=system_report_"
    )
    workbook_bytes = response.get_data()
    assert workbook_bytes.startswith(b"PK")
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    sheet = workbook["System Report"]
    values = list(sheet.iter_rows(values_only=True))
    assert ("Dashboard", "Users", 3) in values
    revenue_row = next(row for row in values if row[1] == "Revenue")
    assert revenue_row[2] == 125.5
